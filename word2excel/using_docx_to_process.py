import os
from docx import Document
import openpyxl

# 定义要解析的文件夹路径和Excel文件路径
# folder_path = "/path/to/folder"
folder_path = "D:\\senior2023\\批量爬取pdf\\pdf2excel\\pdf2excel\\word2excel\\aimed_data"
# excel_file = "/path/to/excel.xlsx"
excel_file = "D:\\senior2023\\批量爬取pdf\\pdf2excel\\pdf2excel\\word2excel\\aimed_data\\excel.xlsx"
# 创建Excel工作簿和工作表
workbook = openpyxl.Workbook()
worksheet = workbook.active

# 遍历文件夹中的所有.docx文件
for filename in os.listdir(folder_path):
    if filename.endswith('.docx'):
        # 使用Document类打开文件
        doc = Document(os.path.join(folder_path, filename))
        
        # 处理文件内容，例如提取标题和正文
        for paragraph in doc.paragraphs:
            if paragraph.style.name.startswith('Heading'):
                # 如果是标题，将标题和空行输出到Excel中
                worksheet.append([paragraph.text, ''])
            else:
                # 如果是正文，将正文输出到Excel中
                worksheet.append(['', paragraph.text])
        
        # 添加一个空行分隔不同的文件
        worksheet.append(['', ''])

# 保存Excel文件
workbook.save(excel_file)
